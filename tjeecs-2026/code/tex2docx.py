# -*- coding: utf-8 -*-
"""
Word version of the manuscript (for reading and editing; the journal gets the LaTeX).

    python tex2docx.py ../manuscript/manuscript.tex

Needs the compiled manuscript (.aux next to the .tex) to know the numbers of the
references, figures, tables and footnotes, and the PNG copies of the figures
(build_all.py writes them). Requires pandoc and python-docx.
Highlighted in the Word file: blue - new sentence on heart rate; red - places that need
the authors' input; orange - author lists completed from memory (verify).
"""
import os, re, sys, subprocess, tempfile
from docx import Document
from docx.shared import Pt, RGBColor

COLORS = {'BLUE': RGBColor(0x1F, 0x4E, 0xC8), 'RED': RGBColor(0xC0, 0x00, 0x00), 'ORANGE': RGBColor(0xE0, 0x7A, 0x00)}


def braces(s, i):
    """Return index after the balanced {...} that starts at s[i] == '{'."""
    depth = 0
    for j in range(i, len(s)):
        if s[j] == '{':
            depth += 1
        elif s[j] == '}':
            depth -= 1
            if depth == 0:
                return j + 1
    raise ValueError('unbalanced braces')


def replace_cmd(s, cmd, fn, nargs=1):
    """Replace \\cmd{a}{b}... by fn(a, b, ...), handling nested braces."""
    out, i = [], 0
    pat = '\\' + cmd + '{'
    while True:
        k = s.find(pat, i)
        if k < 0:
            out.append(s[i:]); break
        out.append(s[i:k]); j = k + len(pat) - 1; args = []
        for _ in range(nargs):
            e = braces(s, j); args.append(s[j + 1:e - 1]); j = e
        out.append(fn(*args)); i = j
    return ''.join(out)


def compress(nums):
    nums = sorted(nums); parts, a = [], 0
    while a < len(nums):
        b = a
        while b + 1 < len(nums) and nums[b + 1] == nums[b] + 1:
            b += 1
        parts.append(f'{nums[a]}–{nums[b]}' if b - a >= 2 else ', '.join(map(str, nums[a:b + 1])))
        a = b + 1
    return ', '.join(parts)


def main(tex_path, out_path=None):
    d = os.path.dirname(os.path.abspath(tex_path))
    tex = open(tex_path, encoding='utf-8').read()
    aux = open(tex_path[:-4] + '.aux', encoding='utf-8').read()
    cites = dict(re.findall(r'\\bibcite\{([^}]*)\}\{\{?(\d+)', aux))
    labels = dict(re.findall(r'\\newlabel\{([^}]*)\}\{\{([^}]*)\}', aux))

    title = re.search(r'\\title\{\s*(.*?)\s*\}', tex, re.S).group(1)
    ru = bool(re.search('[А-Яа-яЁё]', title))
    L = dict(abstract='Аннотация', keywords='Ключевые слова', figure='Рисунок', table='Таблица',
             refs='Список литературы', appendix='Приложение') if ru else \
        dict(abstract='Abstract', keywords='Key words', figure='Figure', table='Table',
             refs='References', appendix='Appendix')
    body = tex.split('\\begin{document}', 1)[1].split('\\end{document}')[0]
    for junk in ('\\setcounter{table}{0}', '\\setcounter{figure}{0}',
                 '\\renewcommand{\\thetable}{A\\arabic{table}}', '\\renewcommand{\\thefigure}{A\\arabic{figure}}'):
        body = body.replace(junk, '')
    body = body.replace('\\maketitle', '\\section*{%s}' % title)
    body = re.sub(r'\\begin\{abstract\}', lambda m: '\\textbf{%s:} ' % L['abstract'], body)
    body = body.replace('\\end{abstract}', '')
    body = replace_cmd(body, 'keywords', lambda k: '\n\n\\textbf{%s:} ' % L['keywords'] + k + '\n\n')
    # tables
    body = re.sub(r'\\input\{(table\w+)\}', lambda m: open(os.path.join(d, m.group(1) + '.tex'), encoding='utf-8').read(), body)
    body = re.sub(r'\\renewcommand\{\\thetable\}\{[^}]*\}', '', body)
    # nested header cells such as \begin{tabular}[c]{@{}c@{}}HR of dominant\\ state\end{tabular}
    body = re.sub(r'\\begin\{tabular\}\[c\]\{@\{\}c@\{\}\}(.*?)\\end\{tabular\}',
                  lambda m: m.group(1).replace('\\\\', ' '), body, flags=re.S)
    # numbers of references, figures, tables, equations, footnotes
    body = replace_cmd(body, 'cite', lambda ks: '[' + compress([int(cites[k]) for k in ks.split(',')]) + ']')
    body = re.sub(r'\\textsuperscript\{\\ref\{([^}]*)\}\}', lambda m: '\\textsuperscript{%s}' % labels.get(m.group(1), '?'), body)
    body = replace_cmd(body, 'ref', lambda k: labels.get(k, '?'))
    body = re.sub(r'\\label\{[^}]*\}', '', body)
    # captions get their number in front
    counters = {'figure': 0, 'table': 0}
    prefix = ['']

    def env(m):
        kind, content = m.group(1), m.group(2)
        counters[kind] += 1
        name = L['figure'] if kind == 'figure' else L['table']
        content = replace_cmd(content, 'caption', lambda c: '\n\n\\textbf{%s %s%d.} %s\n\n' % (name, prefix[0], counters[kind], c))
        content = re.sub(r'\\includegraphics\[[^]]*\]\{([^}]*)\}',
                         lambda g: '\\includegraphics{%s}' % os.path.join(d, g.group(1) + '.png'), content)
        return content
    envpat = r'\\begin\{(figure|table)\}(?:\[[^]]*\])?(.*?)\\end\{\1\}'
    app_head = '\\section*{%s}' % L['appendix']
    if app_head in body:                                # appendix: own numbering A1, A2, ...
        main_, app_ = body.split(app_head, 1)
        main_ = re.sub(envpat, env, main_, flags=re.S)
        counters.update(figure=0, table=0); prefix[0] = 'A'
        for junk in ('\\setcounter{table}{0}', '\\setcounter{figure}{0}',
                     '\\renewcommand{\\thetable}{A\\arabic{table}}', '\\renewcommand{\\thefigure}{A\\arabic{figure}}'):
            app_ = app_.replace(junk, '')
        app_ = re.sub(envpat, env, app_, flags=re.S)
        body = main_ + app_head + app_
    else:
        body = re.sub(envpat, env, body, flags=re.S)
    # highlighting
    body = body.replace('\\chkm{}', '\\chkm').replace('\\chk{}', '\\chk')
    body = body.replace('\\chkm', ' ⟦ORANGE⟧[verify]⟦/⟧').replace('\\chk', ' ⟦RED⟧[add authors/check]⟦/⟧')
    body = replace_cmd(body, 'textcolor', lambda c, t: f'⟦{c.upper()}⟧{t}⟦/⟧', nargs=2)
    # bibliography -> plain numbered paragraphs
    def bib(m):
        items = re.split(r'\\bibitem\{([^}]*)\}', m.group(1))[1:]
        return '\\section*{%s}\n\n' % L['refs'] + '\n\n'.join(f'[{cites[k]}] {t.strip()}' for k, t in zip(items[::2], items[1::2]))
    body = re.sub(r'\\begin\{thebibliography\}\{[^}]*\}(.*?)\\end\{thebibliography\}', bib, body, flags=re.S)
    body = re.sub(r'\\vs\{[^}]*\}', '', body)
    body = body.replace('\\begin{center}', '').replace('\\end{center}', '').replace('\\small', '')

    src = '\\documentclass{article}\n\\usepackage{amsmath}\n\\begin{document}\n' + body + '\n\\end{document}\n'
    tmp = tempfile.mkdtemp()
    open(os.path.join(tmp, 'm.tex'), 'w', encoding='utf-8').write(src)
    out_path = out_path or tex_path[:-4] + '.docx'
    subprocess.run(['pandoc', os.path.join(tmp, 'm.tex'), '-f', 'latex', '-t', 'docx', '-o', out_path], check=True)

    # fonts and colours
    doc = Document(out_path)
    for st in doc.styles:
        try:
            if st.type == 1 and st.font is not None:      # paragraph styles
                st.font.name = 'Times New Roman'
        except Exception:
            pass
    for name in ('Default Paragraph Font',):
        pass
    doc.styles['Normal'].font.size = Pt(11)
    W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    import copy
    mark = re.compile(r'(⟦(?:BLUE|RED|ORANGE)⟧|⟦/⟧)')

    def recolor(par):
        """Split runs at the markers and colour the text between them, in place
        (formulas and other elements of the paragraph keep their positions)."""
        if '⟦' not in ''.join(t.text or '' for t in par._p.iter(W + 't')):
            return
        colour = None
        for r in list(par._p.iter(W + 'r')):
            ts = r.findall(W + 't')
            if not ts:
                continue
            text = ''.join(t.text or '' for t in ts)
            pieces = mark.split(text)
            if len(pieces) == 1 and colour is None:
                continue
            parent = r.getparent(); idx = parent.index(r)
            for piece in pieces:
                if piece.startswith('⟦/'):
                    colour = None; continue
                if piece.startswith('⟦'):
                    colour = piece[1:-1]; continue
                if not piece:
                    continue
                nr = copy.deepcopy(r)
                for t in nr.findall(W + 't'):
                    nr.remove(t)
                t = nr.makeelement(W + 't', {}); t.text = piece
                t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); nr.append(t)
                if colour:
                    rpr = nr.find(W + 'rPr')
                    if rpr is None:
                        rpr = nr.makeelement(W + 'rPr', {}); nr.insert(0, rpr)
                    c = rpr.makeelement(W + 'color', {W + 'val': str(COLORS[colour])}); rpr.append(c)
                parent.insert(idx, nr); idx += 1
            parent.remove(r)

    for p in doc.paragraphs:
        recolor(p)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    recolor(p)
    doc.save(out_path)
    print('Word version:', out_path)


if __name__ == '__main__':
    main(*sys.argv[1:])
